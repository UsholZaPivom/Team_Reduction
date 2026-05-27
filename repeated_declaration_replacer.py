from __future__ import annotations

from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Any
import re

import pandas as pd
from docx import Document


@dataclass
class SafeRule:
    abbreviation: str
    long_form: str
    normalized_long_form: str
    match_count_in_csv: int


@dataclass
class ReplacementEvent:
    abbreviation: str
    long_form: str
    normalized_long_form: str
    source_type: str
    source_index: str
    matched_text: str
    replacement_text: str
    action: str
    occurrence_number_global: int
    comment: str


class RepeatedDeclarationReplacer:
    def __init__(self) -> None:
        self.code_markers = {
            "var", "let", "const", "function", "class", "policy", "key",
            "textblock", "conflict", "return", "select", "insert", "update", "delete"
        }
        self.bad_action_starts = {
            "открыт", "открыта", "открыто", "открытые", "создан", "создана",
            "создано", "создание", "удаление", "удален", "удалена", "получение",
            "внесение", "очистка", "нарушение", "исполнение", "эксплуатация",
            "переполнение", "регистрация", "подбор", "смена", "отключение",
            "исследование", "обход", "разработка", "проектирование"
        }

        self.bad_auto_replace_starts = {
            "области", "область", "сфере", "рамках", "части", "числе",
            "защищенности", "защищённости", "разработанной", "разработанного",
            "обеспечивающих", "обеспечивающие", "программных", "программного",
            "информационным", "информационных", "доступных", "неразмеченных",
            "различных", "разным", "политики", "политик", "средств", "модулей",
            "компонентов", "ресурсам", "ресурсов", "пользователями"
        }

        self.bad_auto_replace_endings = {
            "системы", "документа", "проекта", "работы", "информации",
            "безопасности", "доступа", "записи", "записей"
        }

        self.context_tail_suffixes = (
            "ого", "его", "ой", "ей", "ых", "их", "ым", "им", "ыми", "ими", "ему"
        )

    @staticmethod
    def _clean_text(value: Any) -> str:
        if value is None:
            return ""
        text = str(value).strip()
        return " ".join(text.split()) if text else ""

    @classmethod
    def _normalize_long_form(cls, value: Any) -> str:
        return cls._clean_text(value).lower()

    @staticmethod
    def _extract_words(text: str) -> list[str]:
        return re.findall(r"[A-Za-zА-Яа-яЁё0-9-]+", str(text))

    @staticmethod
    def _make_flexible_whitespace_pattern(text: str) -> str:
        parts = [re.escape(part) for part in str(text).split()]
        return r"\s+".join(parts)

    @staticmethod
    def _spans_overlap(a: tuple[int, int], b: tuple[int, int]) -> bool:
        return max(a[0], b[0]) < min(a[1], b[1])

    def _looks_like_code_or_bad_term(self, text: str) -> bool:
        text = self._clean_text(text)
        if not text:
            return True
        lowered = text.lower()
        words = [w.lower() for w in self._extract_words(text)]
        if words and words[0] in self.bad_action_starts:
            return True
        if any(w in self.code_markers for w in words):
            return True
        if re.search(r"[a-z]+[A-Z][A-Za-z]*", text):
            return True
        if any(symbol in text for symbol in ["=", "{", "}", "<", ">", "\\", "/", "_", "::", "->"]):
            return True
        return False

    def _looks_like_context_tail_for_replacement(self, text: str) -> bool:
        words = [w.lower() for w in self._extract_words(text)]
        if len(words) < 2:
            return True

        first = words[0]
        last = words[-1]

        if first in self.bad_auto_replace_starts:
            return True

        safe_starts = {
            "управление", "система", "модель", "модуль", "механизм", "принцип",
            "контроль", "аудит", "идентификация", "аутентификация",
            "авторизация", "разграничение", "централизованное", "жизненный",
            "цикл", "учетная", "учётная", "учетные", "учётные", "информационная",
            "автоматизированная", "защищенная", "защищённая"
        }
        if first not in safe_starts and any(first.endswith(suffix) for suffix in self.context_tail_suffixes):
            return True

        if len(words) <= 3 and (first in self.bad_auto_replace_starts or last in self.bad_auto_replace_endings):
            return True

        if any(token in {"рисунок", "таблица", "приложение", "раздел", "страница"} for token in words):
            return True

        return False

    def _is_safe_for_auto_replacement(self, abbreviation: str, long_form: str) -> bool:
        abbreviation = self._clean_text(abbreviation).upper()
        long_form = self._clean_text(long_form)
        words = self._extract_words(long_form)

        if not abbreviation or not long_form:
            return False
        if len(words) < 3:
            return False
        if len(words) > 6:
            return False
        if len(abbreviation.replace(" ", "")) < 3:
            return False
        if self._looks_like_code_or_bad_term(long_form):
            return False
        if self._looks_like_context_tail_for_replacement(long_form):
            return False
        return True

    def _should_skip_paragraph_for_replacement(self, paragraph, source_type: str, source_index: str) -> bool:
        if source_type == "front_matter":
            return True

        text = self._clean_text(getattr(paragraph, "text", ""))
        if not text:
            return True

        lowered = text.lower()
        if lowered in {"содержание", "заключение", "список использованных источников"}:
            return True
        if lowered.startswith("рисунок ") or lowered.startswith("таблица "):
            return True

        return False

    @classmethod
    def _build_exact_declaration_regex(cls, long_form: str, abbreviation: str) -> re.Pattern:
        long_pattern = cls._make_flexible_whitespace_pattern(long_form)
        abbr_pattern = cls._make_flexible_whitespace_pattern(abbreviation)
        pattern = rf"(?P<matched>(?P<long>{long_pattern})\s*\(\s*(?:далее\s*[–—-]\s*)?(?P<abbr>{abbr_pattern})\s*\))"
        return re.compile(pattern, flags=re.IGNORECASE)

    @classmethod
    def _build_plain_long_form_regex(cls, long_form: str) -> re.Pattern:
        long_pattern = cls._make_flexible_whitespace_pattern(long_form)
        return re.compile(rf"(?<![A-Za-zА-Яа-яЁё0-9-])(?P<matched>{long_pattern})(?![A-Za-zА-Яа-яЁё0-9-])", flags=re.IGNORECASE)

    def _load_safe_rules(self, existing_abbreviations_csv: Path) -> list[SafeRule]:
        df = pd.read_csv(existing_abbreviations_csv, encoding="utf-8-sig")
        if df.empty:
            return []
        if "abbreviation" not in df.columns:
            raise ValueError("Во входном CSV отсутствует столбец abbreviation.")
        possible_long_columns = [col for col in ["long_form", "matched_term", "term"] if col in df.columns]
        if not possible_long_columns:
            raise ValueError("Во входном CSV отсутствуют long_form, matched_term и term.")

        rules_map: dict[tuple[str, str], SafeRule] = {}
        for _, row in df.iterrows():
            abbreviation = self._clean_text(row.get("abbreviation", "")).upper()
            if not abbreviation:
                continue
            long_form = ""
            for col in possible_long_columns:
                value = self._clean_text(row.get(col, ""))
                if value:
                    long_form = value
                    break
            words = self._extract_words(long_form)
            if not long_form or len(words) < 2:
                continue
            if self._looks_like_code_or_bad_term(long_form):
                continue
            if not self._is_safe_for_auto_replacement(abbreviation, long_form):
                continue
            normalized = self._normalize_long_form(long_form)
            key = (abbreviation, normalized)
            if key not in rules_map:
                rules_map[key] = SafeRule(abbreviation=abbreviation, long_form=long_form, normalized_long_form=normalized, match_count_in_csv=1)
            else:
                rules_map[key].match_count_in_csv += 1

        return sorted(rules_map.values(), key=lambda x: (-len(x.long_form), x.abbreviation, x.long_form))

    def _iter_paragraphs(self, document: Document) -> list[dict[str, Any]]:
        containers = []
        in_front_matter = True

        for i, paragraph in enumerate(document.paragraphs):
            text = self._clean_text(paragraph.text)

            if re.match(r"^1(?:\s|\.|$)", text):
                in_front_matter = False

            source_type = "front_matter" if in_front_matter else "paragraph"
            containers.append({"source_type": source_type, "source_index": str(i), "object": paragraph})

        table_cell_index = 0
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        containers.append({"source_type": "table_cell", "source_index": str(table_cell_index), "object": paragraph})
                        table_cell_index += 1
        return containers

    def _find_matches_in_run(self, text: str, rule: SafeRule) -> list[dict[str, Any]]:
        matches: list[dict[str, Any]] = []
        used: list[tuple[int, int]] = []

        for match in self._build_exact_declaration_regex(rule.long_form, rule.abbreviation).finditer(text):
            span = match.span("matched")
            matches.append({"start": span[0], "end": span[1], "matched_text": match.group("matched"), "match_type": "declaration"})
            used.append(span)

        for match in self._build_plain_long_form_regex(rule.long_form).finditer(text):
            span = match.span("matched")
            if any(self._spans_overlap(span, item) for item in used):
                continue
            matches.append({"start": span[0], "end": span[1], "matched_text": match.group("matched"), "match_type": "plain_long_form"})
            used.append(span)

        matches.sort(key=lambda item: (item["start"], item["end"]))
        return matches

    def _process_run_text(self, text: str, source_type: str, source_index: str, safe_rules: list[SafeRule], occurrence_counters: dict[tuple[str, str], int], events: list[ReplacementEvent]) -> str:
        updated_text = text
        for rule in safe_rules:
            key = (rule.abbreviation, rule.normalized_long_form)
            matches = self._find_matches_in_run(updated_text, rule)
            if not matches:
                continue

            pieces: list[str] = []
            last_pos = 0
            for info in matches:
                start, end = info["start"], info["end"]
                matched_text = info["matched_text"]
                occurrence_counters[key] = occurrence_counters.get(key, 0) + 1
                occurrence_number = occurrence_counters[key]

                if occurrence_number == 1:
                    replacement_text = matched_text
                    action = "keep_first_usage"
                    comment = "Первое употребление полной формы оставлено без изменений."
                else:
                    replacement_text = rule.abbreviation
                    action = "replace_repeated_usage"
                    comment = "Повторное употребление полной формы заменено на сокращение."

                pieces.append(updated_text[last_pos:start])
                pieces.append(replacement_text)
                last_pos = end
                events.append(ReplacementEvent(
                    abbreviation=rule.abbreviation,
                    long_form=rule.long_form,
                    normalized_long_form=rule.normalized_long_form,
                    source_type=source_type,
                    source_index=source_index,
                    matched_text=matched_text,
                    replacement_text=replacement_text,
                    action=action,
                    occurrence_number_global=occurrence_number,
                    comment=comment,
                ))

            pieces.append(updated_text[last_pos:])
            updated_text = "".join(pieces)
        return updated_text

    def _process_paragraph(self, paragraph, source_type: str, source_index: str, safe_rules: list[SafeRule], occurrence_counters: dict[tuple[str, str], int], events: list[ReplacementEvent]) -> None:
        if self._should_skip_paragraph_for_replacement(paragraph, source_type, source_index):
            return

        for run in paragraph.runs:
            original = run.text
            if not self._clean_text(original):
                continue
            updated = self._process_run_text(original, source_type, source_index, safe_rules, occurrence_counters, events)
            if updated != original:
                run.text = updated

    def _build_report_dataframe(self, events: list[ReplacementEvent]) -> pd.DataFrame:
        columns = ["abbreviation", "long_form", "normalized_long_form", "source_type", "source_index", "matched_text", "replacement_text", "action", "occurrence_number_global", "comment"]
        if not events:
            return pd.DataFrame(columns=columns)
        return pd.DataFrame([asdict(item) for item in events])

    def _build_summary_dataframe(self, output_docx_path: Path, safe_rules_count: int, events: list[ReplacementEvent]) -> pd.DataFrame:
        first_kept = sum(1 for e in events if e.action == "keep_first_usage")
        repeated_replaced = sum(1 for e in events if e.action == "replace_repeated_usage")
        return pd.DataFrame([{
            "output_docx": str(output_docx_path).replace("/", "\\"),
            "safe_abbreviation_rules": safe_rules_count,
            "total_matches_found": len(events),
            "first_usages_kept": first_kept,
            "repeated_usages_replaced": repeated_replaced,
            "notes": "Повторные употребления полных форм заменяются только для безопасных терминов и только внутри отдельных run Word, чтобы не ломать форматирование документа.",
        }])

    def _build_statistics_dataframe(self, events: list[ReplacementEvent]) -> pd.DataFrame:
        if not events:
            return pd.DataFrame(columns=["abbreviation", "long_form", "matches_found", "repeated_usages_replaced"])
        df = pd.DataFrame([asdict(item) for item in events])
        grouped = df.groupby(["abbreviation", "long_form"], as_index=False).agg(
            matches_found=("action", "count"),
            repeated_usages_replaced=("action", lambda x: int((x == "replace_repeated_usage").sum())),
        )
        return grouped.sort_values(by=["abbreviation", "long_form"], ascending=[True, True], kind="stable")

    def _save_dataframes(self, output_dir: Path, output_docx_path: Path, report_df: pd.DataFrame, summary_df: pd.DataFrame, statistics_df: pd.DataFrame) -> dict[str, Path]:
        output_dir.mkdir(parents=True, exist_ok=True)
        report_csv = output_dir / f"{output_docx_path.stem}_replacement_report.csv"
        report_xlsx = output_dir / f"{output_docx_path.stem}_replacement_report.xlsx"
        summary_csv = output_dir / f"{output_docx_path.stem}_replacement_summary.csv"
        summary_xlsx = output_dir / f"{output_docx_path.stem}_replacement_summary.xlsx"
        stats_csv = output_dir / f"{output_docx_path.stem}_replacement_statistics_by_abbreviation.csv"
        stats_xlsx = output_dir / f"{output_docx_path.stem}_replacement_statistics_by_abbreviation.xlsx"

        report_df.to_csv(report_csv, index=False, encoding="utf-8-sig")
        report_df.to_excel(report_xlsx, index=False)
        summary_df.to_csv(summary_csv, index=False, encoding="utf-8-sig")
        summary_df.to_excel(summary_xlsx, index=False)
        statistics_df.to_csv(stats_csv, index=False, encoding="utf-8-sig")
        statistics_df.to_excel(stats_xlsx, index=False)

        return {
            "output_docx": output_docx_path,
            "replacement_report_csv": report_csv,
            "replacement_report_xlsx": report_xlsx,
            "replacement_summary_csv": summary_csv,
            "replacement_summary_xlsx": summary_xlsx,
            "replacement_statistics_csv": stats_csv,
            "replacement_statistics_xlsx": stats_xlsx,
        }

    def run(self, source_docx_path: str | Path, existing_abbreviations_csv: str | Path, output_dir: str | Path) -> dict[str, Path]:
        source_docx_path = Path(source_docx_path)
        existing_abbreviations_csv = Path(existing_abbreviations_csv)
        output_dir = Path(output_dir)

        if not source_docx_path.exists():
            raise FileNotFoundError(f"Файл не найден: {source_docx_path}")
        if not existing_abbreviations_csv.exists():
            raise FileNotFoundError(f"Файл не найден: {existing_abbreviations_csv}")

        safe_rules = self._load_safe_rules(existing_abbreviations_csv)
        document = Document(source_docx_path)
        containers = self._iter_paragraphs(document)
        occurrence_counters: dict[tuple[str, str], int] = {}
        events: list[ReplacementEvent] = []

        for container in containers:
            paragraph = container["object"]
            if not self._clean_text(paragraph.text):
                continue
            self._process_paragraph(
                paragraph=paragraph,
                source_type=container["source_type"],
                source_index=container["source_index"],
                safe_rules=safe_rules,
                occurrence_counters=occurrence_counters,
                events=events,
            )

        output_dir.mkdir(parents=True, exist_ok=True)
        output_docx_path = output_dir / f"{source_docx_path.stem}_replaced_declarations.docx"
        document.save(output_docx_path)

        report_df = self._build_report_dataframe(events)
        statistics_df = self._build_statistics_dataframe(events)
        summary_df = self._build_summary_dataframe(output_docx_path, len(safe_rules), events)
        return self._save_dataframes(output_dir, output_docx_path, report_df, summary_df, statistics_df)
